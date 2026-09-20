from __future__ import annotations

import hashlib
import io
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook


MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 2_000
MAX_SHEETS = 100
MAX_ROWS_PER_SHEET = 100_000
MAX_CELLS = 1_000_000
RECHARGE_RATE = Decimal("0.80")
REQUIRED_HEADERS = {"支付时间", "支付金额", "收款项", "支付单号", "订单状态", "用户名"}
TARGET_ITEMS = {
  "收款项：¥50.00 -": (Decimal("50.00"), "¥50.00 -"),
  "收款项：¥100.00 -": (Decimal("100.00"), "¥100.00 -"),
  "收款项：¥200.00 -": (Decimal("200.00"), "¥200.00 -"),
}


def workbook_digest(content: bytes) -> str:
  return hashlib.sha256(content).hexdigest()


def payment_key(payment_number: str) -> str:
  return hashlib.sha256(payment_number.strip().encode("utf-8")).hexdigest()


def masked_payment_number(payment_number: str) -> str:
  clean = payment_number.strip()
  return f"***{clean[-4:]}" if clean else ""


def _text(value) -> str:
  if value is None:
    return ""
  if isinstance(value, datetime):
    return value.strftime("%Y-%m-%d %H:%M:%S")
  if isinstance(value, date):
    return value.strftime("%Y-%m-%d")
  return str(value).strip()


def _amount(value) -> Decimal | None:
  try:
    return Decimal(_text(value)).quantize(Decimal("0.01"))
  except (InvalidOperation, ValueError):
    return None


def validate_xlsx_archive(content: bytes) -> None:
  if not content:
    raise ValueError("XLSX file is required")
  if len(content) > MAX_WORKBOOK_BYTES:
    raise ValueError("XLSX file exceeds the 10 MiB limit")
  try:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
      entries = archive.infolist()
      if len(entries) > MAX_ARCHIVE_ENTRIES:
        raise ValueError("XLSX archive contains too many entries")
      if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
        raise ValueError("XLSX expanded content exceeds the 50 MiB limit")
      names = {entry.filename for entry in entries}
      if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
        raise ValueError("file is not a valid XLSX workbook")
  except zipfile.BadZipFile as exc:
    raise ValueError("file is not a valid XLSX workbook") from exc


def _result(sheet: str, row: int, status: str, reason: str, **values) -> dict:
  return {"sheet": sheet, "row": row, "status": status, "reason": reason, **values}


def parse_recharge_workbook(content: bytes) -> dict:
  validate_xlsx_archive(content)
  try:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
  except Exception as exc:
    raise ValueError("unable to read XLSX workbook") from exc
  try:
    if len(workbook.worksheets) > MAX_SHEETS:
      raise ValueError(f"XLSX workbook exceeds the {MAX_SHEETS}-sheet limit")
    records: list[dict] = []
    target_sheets = 0
    ignored_sheets = 0
    for worksheet in workbook.worksheets:
      # Some payment exports incorrectly declare dimension A1 despite containing many rows.
      worksheet.reset_dimensions()
      rows = []
      cell_count = 0
      for row_number, values in enumerate(worksheet.iter_rows(values_only=True), 1):
        if row_number > MAX_ROWS_PER_SHEET:
          raise ValueError(f"sheet {worksheet.title} exceeds the {MAX_ROWS_PER_SHEET}-row limit")
        cell_count += len(values)
        if cell_count > MAX_CELLS:
          raise ValueError(f"sheet {worksheet.title} exceeds the {MAX_CELLS}-cell limit")
        rows.append((row_number, tuple(values)))

      header_index = None
      headers: dict[str, int] = {}
      for index, (_, values) in enumerate(rows):
        normalized = [_text(value) for value in values]
        if REQUIRED_HEADERS.issubset(set(normalized)):
          header_index = index
          headers = {name: normalized.index(name) for name in REQUIRED_HEADERS}
          break
      metadata = {_text(value) for _, values in rows[: header_index if header_index is not None else len(rows)] for value in values}
      target = next((definition for label, definition in TARGET_ITEMS.items() if label in metadata), None)
      if not target:
        ignored_sheets += 1
        continue
      target_sheets += 1
      denomination, row_item = target
      if header_index is None:
        records.append(_result(worksheet.title, 0, "invalid", "required table headers were not found", amountCny=format(denomination, ".2f")))
        continue

      for row_number, values in rows[header_index + 1 :]:
        if not any(_text(value) for value in values):
          continue
        value = lambda name: values[headers[name]] if headers[name] < len(values) else None
        username = _text(value("用户名"))
        paid_at = _text(value("支付时间"))
        payment_number = _text(value("支付单号"))
        paid_amount = _amount(value("支付金额"))
        common = {
          "username": username,
          "paidAt": paid_at,
          "amountCny": format(denomination, ".2f"),
          "creditCny": format((denomination * RECHARGE_RATE).quantize(Decimal("0.01")), ".2f"),
          "paymentNumber": payment_number,
          "paymentRef": masked_payment_number(payment_number),
        }
        if _text(value("订单状态")) != "支付成功":
          records.append(_result(worksheet.title, row_number, "invalid", "order status is not 支付成功", **common))
        elif paid_amount != denomination:
          records.append(_result(worksheet.title, row_number, "invalid", "payment amount does not match the sheet denomination", **common))
        elif _text(value("收款项")) != row_item:
          records.append(_result(worksheet.title, row_number, "invalid", "row payment item does not match the sheet denomination", **common))
        elif not payment_number:
          records.append(_result(worksheet.title, row_number, "invalid", "支付单号 is required", **common))
        elif not username:
          records.append(_result(worksheet.title, row_number, "invalid", "用户名 is required", **common))
        elif not paid_at:
          records.append(_result(worksheet.title, row_number, "invalid", "支付时间 is required", **common))
        else:
          records.append(_result(worksheet.title, row_number, "candidate", "", **common))
    return {
      "digest": workbook_digest(content),
      "sheetCount": len(workbook.worksheets),
      "targetSheetCount": target_sheets,
      "ignoredSheetCount": ignored_sheets,
      "records": records,
    }
  finally:
    workbook.close()
